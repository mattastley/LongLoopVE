"""Review artifacts, published together only after tune readback succeeds."""

import json
import os
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from longloopve._maxxecu.common import digest
from longloopve.tune import write_tables


def _append(sheet, row):
    sheet.append(row)
    # Treat source metadata as text, even if it begins with an Excel formula prefix.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def workbook(result, path):
    book = Workbook()
    overview = book.active
    overview.title = "Summary"
    for key in (
        "nickname",
        "car",
        "epoch",
        "archived",
        "created_at",
        "reference_tune",
        "latest_tune",
    ):
        _append(overview, [key, result[key]])
    _append(overview, ["Interpolation/layout evidence", result["mapping"]["layout_evidence"]])
    _append(
        overview,
        ["Method", "Equal-sample least squares; minimum change from calibration reference"],
    )
    _append(overview, ["Review", "Human review required before loading exported tune into ECU"])
    for key, value in result["settings"].items():
        _append(overview, [key, value])
    for bank, data in result["banks"].items():
        for key in ("sample_count", "rank", "condition_estimate", "fit_rmse", "export_rmse"):
            _append(overview, [f"Bank {bank}: {key}", data[key]])
        for warning in data["warnings"]:
            _append(overview, [f"Bank {bank}: warning", warning])
        for field, label in (
            ("estimate", "Estimate"),
            ("exported", "Export VE"),
            ("base", "Base VE"),
            ("change_percent", "Change pct"),
            ("counts", "Counts"),
            ("weight_sums", "Weights"),
            ("log_counts", "Logs"),
            ("eligible", "Eligible"),
        ):
            sheet = book.create_sheet(f"{bank} {label}")
            _append(sheet, ["ETPS % / RPM", *data["rpm"]])
            for index in reversed(range(len(data["etps"]))):
                _append(sheet, [data["etps"][index], *data[field][index]])
            sheet.freeze_panes = "B2"
            sheet.column_dimensions["A"].width = 18
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(horizontal="center")
                    if cell.row == 1 or cell.column == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="E9EEF4")
                    elif field not in {"counts", "log_counts", "eligible"}:
                        cell.number_format = "0.00"
            if field == "change_percent":
                scale = result["settings"]["color_scale"]
                sheet.conditional_formatting.add(
                    f"B2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}",
                    ColorScaleRule(
                        start_type="num",
                        start_value=-scale,
                        start_color="70BF94",
                        mid_type="num",
                        mid_value=0,
                        mid_color="FFFFFF",
                        end_type="num",
                        end_value=scale,
                        end_color="E89191",
                    ),
                )
    logs = book.create_sheet("Log provenance")
    _append(
        logs,
        [
            "SHA256",
            "Source",
            "Name/Car",
            "Decoder",
            "Device validation",
            "Recording tune",
            "Warnings",
        ],
    )
    filters = book.create_sheet("Filters")
    _append(filters, ["Log", "Bank", "Total", "Accepted", "Rejected", "Reason", "Count"])
    for log in result["logs"]:
        _append(
            logs,
            [
                log["log"],
                log["source_name"],
                json.dumps(log["identity"]),
                json.dumps(log["decoder"]),
                log["device_validation"],
                log["recording_tune"],
                json.dumps(log["warnings"]),
            ],
        )
        for bank, summary in log["summary"].items():
            for reason, count in summary["reasons"].items():
                _append(
                    filters,
                    [
                        log["log"],
                        bank,
                        summary["total"],
                        summary["accepted"],
                        summary["rejected"],
                        reason,
                        count,
                    ],
                )
            for warning in summary["warnings"]:
                _append(overview, [f"{log['source_name']} bank {bank}", warning])
    _append(filters, ["Note", "Reason counts overlap; one observation can fail multiple rules"])
    overview.column_dimensions["A"].width = 35
    overview.column_dimensions["B"].width = 100
    overview.freeze_panes = "B2"
    book.save(path)


def png(result, path):
    # Keep headless CLI caches writable without writing into a user's home directory.
    cache = Path(tempfile.gettempdir()) / f"longloopve-cache-{os.getuid()}"
    cache.mkdir(exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(cache / "matplotlib"))
    os.environ.setdefault("XDG_CACHE_HOME", str(cache))
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    from matplotlib.colors import LinearSegmentedColormap, Normalize
    from matplotlib.patches import Rectangle

    a, b = result["banks"]["A"], result["banks"]["B"]
    nx, ny = len(a["rpm"]), len(a["etps"])
    fig, ax = plt.subplots(figsize=(max(8, nx * 0.85), max(5, ny * 0.65 + 2.5)))
    scale = result["settings"]["color_scale"]
    colors = LinearSegmentedColormap.from_list("ve_change", ["#72be94", "#ffffff", "#e78b8b"])
    norm = Normalize(-scale, scale, clip=True)
    for row, index in enumerate(reversed(range(ny))):
        for column in range(nx):
            for offset, bank, data in ((0, "A", a), (0.5, "B", b)):
                count = data["counts"][index][column]
                eligible = data["eligible"][index][column]
                change = data["change_percent"][index][column]
                color = colors(norm(change)) if eligible and change is not None else "#edf0f4"
                ax.add_patch(
                    Rectangle(
                        (column, row + offset),
                        1,
                        0.5,
                        facecolor=color,
                        edgecolor="white",
                        linewidth=0.6,
                        hatch="//" if count and not eligible else None,
                    )
                )
                label = f"{bank} {change:+.1f}" if eligible and change is not None else f"{bank} —"
                ax.text(
                    column + 0.5,
                    row + (0.2 if bank == "A" else 0.8),
                    label,
                    ha="center",
                    va="center",
                    fontsize=8,
                    color="#293f53" if eligible else "#8995a2",
                )
            ca, cb = a["counts"][index][column], b["counts"][index][column]
            count_label = f"N {ca:,}" if ca == cb else f"N A {ca:,}\nB {cb:,}"
            ax.text(
                column + 0.5,
                row + 0.5,
                count_label,
                ha="center",
                va="center",
                fontsize=5.5,
                linespacing=1,
                color="#667483",
                bbox={"facecolor": "white", "edgecolor": "none", "pad": 0.2},
            )
    ax.set(xlim=(0, nx), ylim=(ny, 0), ylabel="Actual throttle body position (%)")
    ax.set_xticks(np.arange(nx) + 0.5, [f"{v:g}" for v in a["rpm"]])
    ax.set_yticks(
        np.arange(ny) + 0.5,
        [f"{v:.1f}".rstrip("0").rstrip(".") for v in reversed(a["etps"])],
    )
    ax.xaxis.tick_top()
    ax.xaxis.set_label_position("top")
    ax.set_xlabel("RPM")
    ax.tick_params(length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.suptitle(f"VE Change Report — {result['nickname']} — Bank A/B", fontsize=16, y=0.985)
    fig.text(
        0.5,
        0.93,
        "Upper: Bank A change (%)     Lower: Bank B change (%)     Center: observations",
        ha="center",
        fontsize=9,
    )
    warning_count = sum(len(d["warnings"]) for d in result["banks"].values())
    fig.text(
        0.5,
        0.055,
        f"Red: positive | Green: negative | Gray: no data | Hatched: below minimum "
        f"({result['settings']['min_observations']}) | Color scale: ±{scale:g}%",
        ha="center",
        fontsize=8,
    )
    fig.text(
        0.5,
        0.025,
        f"Changes vs latest base {result['latest_tune'][:12]}; {len(result['logs'])} logs; "
        f"{warning_count} warnings (see workbook). Review before use.",
        ha="center",
        fontsize=8,
    )
    fig.subplots_adjust(top=0.85, bottom=0.12, left=0.09, right=0.99)
    fig.savefig(path, dpi=160)
    plt.close(fig)


def export(result, destination: Path):
    destination = destination.absolute()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.mkdir(exist_ok=False)
    published = False
    try:
        with tempfile.TemporaryDirectory(prefix=".ve-export-", dir=destination.parent) as temporary:
            stage = Path(temporary) / "result"
            stage.mkdir()
            original = Path(result["latest_tune_original"]).read_bytes()
            if digest(original) != result["latest_tune"]:
                raise ValueError("Retained latest tune was modified; export aborted")
            tune = write_tables(
                original, result["mapping"], {b: result["banks"][b]["exported"] for b in ("A", "B")}
            )
            (stage / "corrected.MaxxECU-save").write_bytes(tune)
            workbook(result, stage / "ve-tables.xlsx")
            png(result, stage / "ve-change.png")
            (stage / "calculation.json").write_text(
                json.dumps(result, indent=2, allow_nan=False) + "\n"
            )
            os.replace(stage, destination)
            published = True
        return {
            "output": str(destination),
            "files": [
                "ve-tables.xlsx",
                "ve-change.png",
                "corrected.MaxxECU-save",
                "calculation.json",
            ],
            "warnings": {b: result["banks"][b]["warnings"] for b in ("A", "B")},
        }
    finally:
        if not published:
            destination.rmdir()
