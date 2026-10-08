"""Known-answer reconstruction and transactional, independent-bank workflows."""

import json
import sys
from pathlib import Path

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from longloopve.analysis import (
    BANK_ROLES,
    COMMON,
    FILTERS,
    Channel,
    interpolation,
    observations,
    reconstruct,
    settings,
)
from longloopve.cli import main
from longloopve.library import Library, identity
from longloopve.reports import export
from longloopve.tune import read_tables, write_tables


def mapping(order="rpm_fast", scale=1):
    return {
        "interpolation": "bilinear_clamped",
        "layout_evidence": "Synthetic known-answer fixture",
        "banks": {
            b: {
                "table": f"VE {b}",
                "rpm_axis": "X",
                "etps_axis": "Y",
                "rpm_scale": 1,
                "etps_scale": 1,
                "value_scale": scale,
                "order": order,
            }
            for b in ("A", "B")
        },
    }


def tune(values=None, order="rpm_fast", encoding="utf-8", rpm="1000,2000", scale=1):
    values = values if values is not None else [[50, 70], [80, 100]]
    data = np.asarray(values) / scale
    if order == "etps_fast":
        data = data.T
    text = f'<?xml version="1.0" encoding="{encoding}"?><MaxxECUSettingsFile>'
    text += "<Fileinfo><Softwareversion>1.159</Softwareversion></Fileinfo>"
    text += (
        '<!-- retain this comment --><FutureThing attr="001"><Nested>keep me</Nested></FutureThing>'
    )
    text += '<ECUSettingsItem name="Math">VE * LTT * Lambda</ECUSettingsItem>'
    for bank in ("A", "B"):
        text += (
            f'<DynamicTable name="VE {bank}" future="unchanged">'
            f'<DynamicTableAxis name="X" AxisCellCount="2">{rpm}</DynamicTableAxis>'
            '<DynamicTableAxis name="Y" AxisCellCount="2">0,100</DynamicTableAxis>'
            '<DynamicTableData XSize="2" YSize="2">'
            + ",\n".join(f" {v:g} " for v in data.ravel())
            + "</DynamicTableData></DynamicTable>"
        )
    return (text + "</MaxxECUSettingsFile>").encode(encoding)


def profile():
    specs = {}
    for role in COMMON:
        units = "rpm" if role == "engine_speed" else "%"
        if role in ("coolant_temperature", "intake_air_temperature"):
            units = "C"
        specs[role] = {"channel": role, "units": units}
        if role in FILTERS:
            specs[role]["inactive"] = 0
    for bank in ("a", "b"):
        for role in BANK_ROLES:
            role = role.format(bank=bank)
            specs[role] = {"channel": role, "units": "%"}
    return {
        "channels": specs,
        "sanity": {
            "factors": [
                {"role": "long_term_trim_bank_{bank}", "convention": "offset_percent"},
                {"role": "lambda_correction_bank_{bank}", "convention": "multiplier_percent"},
            ]
        },
    }


def channels(times=None):
    times = np.asarray(times if times is not None else [0, 10, 20, 30])
    values = {role: np.zeros(len(times)) for role in profile()["channels"]}
    values["engine_speed"] = np.resize([1000, 2000, 1000, 2000], len(times))
    values["throttle_body_position"] = np.resize([0, 0, 100, 100], len(times))
    values["coolant_temperature"][:] = 80
    values["intake_air_temperature"][:] = 40
    for b in ("a", "b"):
        values[f"suggested_ve_bank_{b}"] = np.resize([55, 77, 88, 110], len(times))
        values[f"current_ve_bank_{b}"] = values[f"suggested_ve_bank_{b}"] / 1.1
        values[f"lambda_correction_bank_{b}"][:] = 110
    return {role: Channel(times, v) for role, v in values.items()}


def record(rpm, etps, values):
    return {
        "rpm": np.asarray(rpm),
        "etps": np.asarray(etps),
        "values": np.asarray(values),
        "summary": {"reasons": {}, "warnings": []},
    }


def grid(values=None):
    return {
        "rpm": [1000, 2000],
        "etps": [0, 100],
        "values": values if values is not None else [[50, 70], [80, 100]],
    }


def fake_ingest(source, output, **kwargs):
    output.mkdir()
    (output / "channels").mkdir()
    source_config = json.loads(source.read_bytes())
    cs = channels()
    if source_config.get("bad_b"):
        cs["suggested_ve_bank_b"].values[:] = np.nan
    entries = []
    for i, (role, channel) in enumerate(cs.items()):
        if source_config.get("missing") == role:
            continue
        units = profile()["channels"][role]["units"]
        schema = pa.schema(
            [
                pa.field("timecodes", pa.int64()),
                pa.field(role, pa.float64(), metadata={b"units": units.encode()}),
            ]
        )
        table = pa.Table.from_arrays(
            [pa.array(channel.times, type=pa.int64()), pa.array(channel.values)], schema=schema
        )
        path = f"channels/{i:05d}.parquet"
        pq.write_table(table, output / path)
        entries.append({"name": role, "units": units, "file": path})
    from longloopve._maxxecu.common import digest

    manifest = {
        "metadata": {"Name": "Driver", "Car": "S54"},
        "channels": entries,
        "source": {"sha256": digest(source.read_bytes())},
        "decoder": {"name": "fixture", "version": "1"},
        "device_validation": "synthetic",
        "warnings": [],
    }
    (output / "manifest.json").write_text(json.dumps(manifest))
    return manifest


@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setattr("longloopve.library.run", fake_ingest)
    lib = Library(tmp_path / "library")
    path = tmp_path / "base.MaxxECU-save"
    path.write_bytes(tune())
    car = lib.create("Driver", "S54", path, mapping(), profile())
    return lib, car, path


def test_nonuniform_interpolation_nodes_edges_and_clamping():
    xs, ys = np.array([500, 1100, 2500]), np.array([0, 13, 100])
    cells, weights = interpolation([500, 800, 9999], [0, 6.5, 999], xs, ys)
    np.testing.assert_allclose(weights.sum(axis=1), 1)
    np.testing.assert_allclose(weights[0], [1, 0, 0, 0])
    np.testing.assert_allclose(weights[1], [0.25] * 4)
    assert cells[2, 3] == 8 and weights[2, 3] == 1


def test_reconstructs_known_sloping_table_from_interpolated_samples():
    rng = np.random.default_rng(29)
    rpm, etps = rng.uniform(1000, 2000, 100), rng.uniform(0, 100, 100)
    truth = np.array([[60, 75], [95, 120]])
    cells, weights = interpolation(rpm, etps, np.array([1000, 2000]), np.array([0, 100]))
    targets = (truth.ravel()[cells] * weights).sum(axis=1)
    result = reconstruct(grid(), grid(), [record(rpm, etps, targets)])
    np.testing.assert_allclose(result["estimate"], truth, atol=1e-7)
    assert result["fit_rmse"] < 1e-7 and result["rank"] == 4
    assert result["counts"] == [[100, 100], [100, 100]]
    assert np.sum(result["weight_sums"]) == pytest.approx(100)


def test_minimum_change_solution_and_latest_base_independence():
    prior = grid()
    latest = grid([[90, 90], [90, 90]])
    obs = record([1500], [50], [100])
    first = reconstruct(prior, prior, [obs])
    second = reconstruct(prior, latest, [obs])
    np.testing.assert_allclose(first["estimate"], [[75, 95], [105, 125]])
    np.testing.assert_allclose(first["estimate"], second["estimate"])
    assert second["rank"] == 1
    assert any("Underdetermined" in w for w in second["warnings"])


def test_equal_sample_influence_counts_and_export_threshold():
    result = reconstruct(
        grid(),
        grid([[90, 90], [90, 90]]),
        [record([1000], [0], [60]), record([1000, 1000], [0, 0], [90, 90])],
        minimum=4,
    )
    assert result["estimate"][0][0] == pytest.approx(80)
    assert result["counts"] == [[3, 0], [0, 0]]
    assert result["log_counts"] == [[2, 0], [0, 0]]
    assert result["exported"] == [[90, 90], [90, 90]]
    assert result["export_rmse"] > result["fit_rmse"]


def test_alignment_exact_linear_hold_stale_and_gaps():
    channel = Channel([0, 20, 1000], [0, 20, 1000])
    values = channel.align(np.array([-1, 0, 10, 20, 100, 999, 1000, 1001]), 30, "linear")
    np.testing.assert_allclose(values[[1, 2, 3, 6]], [0, 10, 20, 1000])
    assert np.isnan(values[[0, 4, 5, 7]]).all()
    held = channel.align(np.array([10, 51]), 30)
    assert held[0] == 0 and np.isnan(held[1])
    invalid = Channel([0, 0], [1, 2])
    assert np.isnan(invalid.align(np.array([0]), 30)).all()


@pytest.mark.parametrize("role", FILTERS)
def test_cut_and_enrichment_filter_rejects_each_bank(role):
    cs = channels()
    cs[role].values[1] = 1
    result = observations(cs, profile(), settings())
    assert result["A"]["summary"]["accepted"] == 3
    assert result["B"]["summary"]["reasons"][role] == 1


def test_temperature_boundaries_and_bank_independence():
    cs = channels()
    cs["coolant_temperature"].values = np.array([165, 166, 166, 166])
    cs["intake_air_temperature"].values = np.array([100, 130, 131, 100])
    p = profile()
    p["channels"]["coolant_temperature"]["units"] = "F"
    p["channels"]["intake_air_temperature"]["units"] = "F"
    cs["suggested_ve_bank_b"].values[1] = np.nan
    result = observations(cs, p, settings())
    assert result["A"]["summary"]["accepted"] == 2
    assert result["B"]["summary"]["accepted"] == 1
    assert result["A"]["summary"]["warnings"] == []


def test_sanity_discrepancy_warns_without_rejection():
    cs = channels()
    cs["suggested_ve_bank_a"].values *= 1.2
    result = observations(cs, profile(), settings())
    assert result["A"]["summary"]["accepted"] == 4
    assert "discrepancy" in result["A"]["summary"]["warnings"][0]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-8-sig", "cp1252"])
@pytest.mark.parametrize("order", ["rpm_fast", "etps_fast"])
def test_tune_text_only_edits_scaling_encoding_unknown_elements(encoding, order):
    data = tune(order=order, encoding=encoding, scale=0.1)
    updates = {"A": [[51, 70], [80, 100]], "B": [[50, 70], [80, 101]]}
    output = write_tables(data, mapping(order, 0.1), updates)
    parsed = read_tables(output, mapping(order, 0.1))
    assert parsed["A"]["values"] == updates["A"]
    text = output.decode(encoding)
    assert "<!-- retain this comment -->" in text
    assert '<FutureThing attr="001"><Nested>keep me</Nested></FutureThing>' in text
    # Only two numeric substrings change, even for inconsistent XML encoding declarations.
    assert text.replace(" 510 ", " 500 ").replace(" 1010 ", " 1000 ") == data.decode(encoding)


def test_tune_rejects_ambiguous_mapping_and_protected_content():
    wrong = mapping()
    wrong["banks"]["B"]["table"] = "VE A"
    with pytest.raises(ValueError, match="distinct"):
        read_tables(tune(), wrong)
    protected = tune().replace(
        b"</MaxxECUSettingsFile>", b"<EncryptedSettings>x</EncryptedSettings></MaxxECUSettingsFile>"
    )
    with pytest.raises(ValueError, match="partial"):
        read_tables(protected, mapping())


def test_library_dedup_preserve_reset_reassignment_and_grid_rejection(library, tmp_path):
    lib, car, base = library
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    imported = lib.import_log(source)
    assert imported["car"] == car["id"]
    assert lib.import_log(source)["status"] == "already_imported"
    before = lib.calculate(car["id"])
    update = tmp_path / "updated.MaxxECU-save"
    update.write_bytes(tune([[90, 90], [90, 90]]))
    lib.add_tune(car["id"], update, "preserve")
    after = lib.calculate(car["id"])
    assert before["banks"]["A"]["estimate"] == after["banks"]["A"]["estimate"]
    assert after["banks"]["A"]["base"] == [[90, 90], [90, 90]]
    old_epoch = before["epoch"]
    lib.add_tune(car["id"], update, "reset")
    assert lib.calculate(car["id"])["banks"]["A"]["sample_count"] == 0
    assert lib.calculate(car["id"], epoch_id=old_epoch)["banks"]["A"]["sample_count"] == 4
    assert lib.calculate(car["id"], epoch_id=old_epoch)["archived"]
    lib.assign_log(imported["log"], car["id"])
    assert lib.calculate(car["id"])["banks"]["A"]["sample_count"] == 4
    assert not lib.calculate(car["id"], epoch_id=old_epoch)["logs"]
    update.write_bytes(tune(rpm="1000,2500"))
    with pytest.raises(ValueError, match="breakpoints"):
        lib.add_tune(car["id"], update, "preserve")
    assert lib.calculate(car["id"])["banks"]["A"]["base"] == [[90, 90], [90, 90]]


def test_missing_required_channel_fails_without_assignment(library, tmp_path):
    lib, car, _ = library
    source = tmp_path / "missing.xrk"
    source.write_text('{"missing":"knock_retard"}')
    with pytest.raises(ValueError, match="absent"):
        lib.import_log(source)
    assert lib.load()["logs"] == {}
    source.write_text('{"bad_b":true}')
    lib.import_log(source)
    result = lib.calculate(car["id"])
    assert result["banks"]["A"]["sample_count"] == 4
    assert result["banks"]["B"]["sample_count"] == 0
    assert result["banks"]["B"]["exported"] == grid()["values"]


def test_filter_recalculation_and_recording_tune_sanity(library, tmp_path):
    lib, car, _ = library
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    epoch = car["epochs"][car["active_epoch"]]
    lib.import_log(source, recording_tune=epoch["latest_tune"])
    assert lib.calculate(car["id"])["banks"]["A"]["sample_count"] == 4
    filtered = lib.calculate(car["id"], {"clt_min_f": 180})
    assert filtered["banks"]["A"]["sample_count"] == 0
    assert all("Recording-time tune unknown" not in w for w in filtered["banks"]["A"]["warnings"])


def test_transaction_lock_and_failure_preserves_index(library):
    lib, _, _ = library
    before = lib.index_path.read_bytes()
    with pytest.raises(RuntimeError):
        with lib.writer() as index:
            index["cars"].clear()
            with pytest.raises(ValueError, match="lock"):
                with lib.writer():
                    pass
            raise RuntimeError()
    assert lib.index_path.read_bytes() == before
    assert not (lib.root / ".writer.lock").exists()
    with pytest.raises(ValueError, match="escapes"):
        lib.path("../outside")


def test_exports_roundtrip_workbook_png_and_no_overwrite(library, tmp_path):
    lib, car, _ = library
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    lib.import_log(source)
    target = tmp_path / "review"
    result = lib.export(car["id"], target)
    assert set(result["files"]) == {p.name for p in target.iterdir()}
    assert (target / "ve-change.png").read_bytes().startswith(b"\x89PNG")
    book = load_workbook(target / "ve-tables.xlsx")
    assert book["A Export VE"]["B3"].value == pytest.approx(55)
    assert book["A Counts"]["B3"].value == 1
    assert read_tables((target / "corrected.MaxxECU-save").read_bytes(), mapping())["A"]["values"][
        0
    ][0] == pytest.approx(55)
    with pytest.raises(FileExistsError):
        lib.export(car["id"], target)


def test_failed_export_leaves_no_partial_bundle(library, tmp_path, monkeypatch):
    lib, car, _ = library

    def fail(*args):
        raise RuntimeError("plot failure")

    monkeypatch.setattr("longloopve.reports.png", fail)
    target = tmp_path / "review"
    with pytest.raises(RuntimeError, match="plot failure"):
        lib.export(car["id"], target)
    assert not target.exists()


def test_metadata_formula_injection_and_case_sensitive_identity(library, tmp_path):
    lib, car, _ = library
    result = lib.calculate(car["id"])
    result["nickname"] = '=HYPERLINK("https://example.com")'
    export(result, tmp_path / "review")
    book = load_workbook(tmp_path / "review" / "ve-tables.xlsx")
    assert book["Summary"]["B1"].data_type == "s"
    assert identity({"Name": "Driver", "Car": "S54"}, {}) == ["Driver", "S54"]
    assert identity({"filename": "Driver_S54"}, {}) == [None, None]


def test_packaged_ingest_helpers_match_both_skills():
    root = Path(__file__).resolve().parents[1]
    for name in ("__init__.py", "tunes.py", "common.py"):
        installed = root / "src" / "longloopve" / "_maxxecu" / name
        for skill in ("maxxecu-tune-ingest", "maxxecu-log-ingest"):
            assert (
                installed.read_bytes()
                == (
                    root / ".agents" / "skills" / skill / "scripts" / "maxxecu_backend" / name
                ).read_bytes()
            )


def test_cli_persistent_commands_and_actionable_failure(tmp_path, monkeypatch, capsys):
    base = tmp_path / "base.MaxxECU-save"
    base.write_bytes(tune())
    m, p = tmp_path / "mapping.json", tmp_path / "profile.json"
    m.write_text(json.dumps(mapping()))
    p.write_text(json.dumps(profile()))

    def invoke(*arguments):
        monkeypatch.setattr(
            sys, "argv", ["longloopve", "--library", str(tmp_path / "library"), *arguments]
        )
        code = main()
        captured = capsys.readouterr()
        return code, captured

    code, output = invoke(
        "car-add",
        "--name",
        "Driver",
        "--identifier",
        "S54",
        "--tune",
        str(base),
        "--mapping",
        str(m),
        "--profile",
        str(p),
    )
    assert code == 0
    car = json.loads(output.out)["id"]
    assert invoke("cars")[0] == 0
    assert invoke("calculate", "--car", car)[0] == 0
    assert invoke("calculate", "--car", car, "--min-observations", "0")[0] == 1
    assert invoke("tune-inspect", str(base))[0] == 0
    code, output = invoke("calculate", "--car", "missing")
    assert code == 1 and "missing" in output.err


def test_unusable_bank_timeline_does_not_block_other_bank():
    cs = channels()
    cs["suggested_ve_bank_b"] = Channel([0, 10, 10, 30], [50, 60, 70, 80])
    result = observations(cs, profile(), settings())
    assert result["A"]["summary"]["accepted"] == 4
    assert result["B"]["summary"]["accepted"] == 0
    assert result["B"]["summary"]["reasons"]["invalid_corrected_timeline"] == 4


def test_zero_base_change_is_unavailable_and_exportable(library, tmp_path):
    lib, car, _ = library
    updated = tmp_path / "zero.MaxxECU-save"
    updated.write_bytes(tune([[0, 70], [80, 100]]))
    lib.add_tune(car["id"], updated, "preserve")
    result = lib.calculate(car["id"])
    assert result["banks"]["A"]["change_percent"][0][0] is None
    lib.export(car["id"], tmp_path / "review")
    assert (tmp_path / "review" / "ve-change.png").exists()


def test_unit_mismatch_is_error_and_automatic_identity_requires_both_fields(library, tmp_path):
    lib, car, _ = library
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    imported = lib.import_log(source)
    folder = lib.path(imported["directory"])
    manifest_path = folder / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["channels"][0]["units"] = "wrong"
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="Unit mismatch"):
        lib.calculate(car["id"])
    assert identity({"Driver": "Driver", "Vehicle": "S54"}, {}) == ["Driver", "S54"]
    assert identity({"Name": "Other", "Driver": "Driver", "Vehicle": "S54"}, {}) == [None, "S54"]
    assert identity({"Car": "S54"}, {}) == [None, "S54"]


def test_library_is_relocatable(library, tmp_path):
    import shutil

    lib, car, _ = library
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    lib.import_log(source)
    target = tmp_path / "moved"
    shutil.copytree(lib.root, target)
    moved = Library(target)
    assert moved.calculate(car["id"])["banks"] == lib.calculate(car["id"])["banks"]
    moved.export(car["id"], tmp_path / "review")


def test_modified_retained_tune_aborts_export(library, tmp_path):
    lib, car, _ = library
    result = lib.calculate(car["id"])
    Path(result["latest_tune_original"]).write_bytes(tune([[90, 90], [90, 90]]))
    with pytest.raises(ValueError, match="modified"):
        lib.export(car["id"], tmp_path / "review")
    assert not (tmp_path / "review").exists()


def test_empty_bank_and_unequal_timestamps_have_independent_counts():
    cs = channels()
    cs["suggested_ve_bank_b"] = Channel([5, 25], [60, 90])
    result = observations(cs, profile(), settings())
    assert result["A"]["summary"]["accepted"] == 4
    assert result["B"]["summary"]["accepted"] == 2
    np.testing.assert_allclose(result["B"]["rpm"], [1500, 1500])
    cs["suggested_ve_bank_b"] = Channel([], [])
    result = observations(cs, profile(), settings())
    assert result["B"]["summary"]["accepted"] == 0


def test_identity_resolution_new_and_conflicting_explicit_assignment(library, tmp_path):
    lib, first, base = library
    other = lib.create("Other", "OtherCar", base, mapping(), profile())
    source = tmp_path / "sample.xrk"
    source.write_text("{}")
    with pytest.raises(ValueError, match="another car"):
        lib.import_log(source, car_id=other["id"])
    assert not lib.load()["logs"]
    # Forced selection does not change the registered identity of the first car.
    assert lib.load()["cars"][first["id"]]["identities"] == [["Driver", "S54"]]


def test_batch_cli_continues_after_failures(library, tmp_path, monkeypatch, capsys):
    lib, _, _ = library
    bad, good = tmp_path / "bad.xrk", tmp_path / "good.xrk"
    bad.write_text('{"missing":"knock_retard"}')
    good.write_text("{}")
    monkeypatch.setattr(
        sys, "argv", ["longloopve", "--library", str(lib.root), "log-add", str(bad), str(good)]
    )
    assert main() == 2
    output = json.loads(capsys.readouterr().out)
    assert len(output["errors"]) == 1 and len(output["files"]) == 1
    assert len(lib.load()["logs"]) == 1


@pytest.mark.parametrize(
    "overrides",
    [
        {"min_observations": 0},
        {"min_observations": 1.5},
        {"freshness_ms": 0},
        {"clt_min_f": float("nan")},
        {"unknown": 3},
        {"color_scale": -1},
    ],
)
def test_invalid_settings_fail_explicitly(overrides):
    with pytest.raises(ValueError):
        settings(overrides)


def test_no_percentage_cap_including_signed_reconstructed_values():
    updates = {"A": [[-20, 300], [80, 100]], "B": [[50, 70], [80, 100]]}
    output = write_tables(tune(), mapping(), updates)
    assert read_tables(output, mapping(), allow_negative=True)["A"]["values"] == updates["A"]


def test_first_log_can_register_new_car_through_resolver(tmp_path, monkeypatch):
    monkeypatch.setattr("longloopve.library.run", fake_ingest)
    lib = Library(tmp_path / "library")
    source, base = tmp_path / "sample.xrk", tmp_path / "base.MaxxECU-save"
    source.write_text("{}")
    base.write_bytes(tune())

    def resolve(metadata, index):
        assert metadata == {"Name": "Driver", "Car": "S54"}
        assert not index["cars"]
        return {
            "name": "Driver",
            "identifier": "S54",
            "tune": base,
            "mapping": mapping(),
            "profile": profile(),
        }

    imported = lib.import_log(source, resolver=resolve)
    assert len(lib.load()["cars"]) == 1
    assert lib.calculate(imported["car"])["banks"]["A"]["sample_count"] == 4


@pytest.mark.parametrize("order", ["rpm_fast", "etps_fast"])
def test_nonsquare_tune_orientation_against_independent_source(order):
    m = mapping(order)
    values = "10,20,30,40,50,60" if order == "rpm_fast" else "10,40,20,50,30,60"
    rpm_axis, etps_axis = ("X", "Y") if order == "rpm_fast" else ("Y", "X")
    xsize, ysize = (3, 2) if order == "rpm_fast" else (2, 3)
    xml = "<MaxxECUSettingsFile>"
    for b in "AB":
        m["banks"][b]["rpm_axis"] = rpm_axis
        m["banks"][b]["etps_axis"] = etps_axis
        xml += (
            f'<DynamicTable name="VE {b}">'
            f'<DynamicTableAxis name="{rpm_axis}" AxisCellCount="3">'
            "1000,1800,3000</DynamicTableAxis>"
            f'<DynamicTableAxis name="{etps_axis}" AxisCellCount="2">0,100</DynamicTableAxis>'
            f'<DynamicTableData XSize="{xsize}" YSize="{ysize}">{values}</DynamicTableData>'
            "</DynamicTable>"
        )
    data = (xml + "</MaxxECUSettingsFile>").encode()
    assert read_tables(data, m)["A"]["values"] == [[10, 20, 30], [40, 50, 60]]
    updates = {b: [[11, 21, 31], [41, 51, 61]] for b in "AB"}
    written = write_tables(data, m, updates)
    assert read_tables(written, m)["A"]["values"] == updates["A"]
